"""
Importador de la planilla de escolaridades.

La planilla real tiene 20.000 filas y esta cifrada; aca se arma una chica con
openpyxl que reproduce cada caso raro que aparecio al analizarla: duplicados
exactos, filas con solo el operador, fechas tipeadas a mano, '#REF!', codigos
de plan con y sin espacio, filas de ALUMNOS desplazadas, empresas mezcladas
con personas, cedulas repetidas y gente que solo aparece en las actas.
"""
import pytest
from datetime import date, datetime

from openpyxl import Workbook
from sqlmodel import select, func

from v2.models.historico import HistoricoAlumno, HistoricoPlan, HistoricoResultado
from v2.scripts.importar_historico import (
    importar, codigo_plan, parsear_fecha, parsear_cedula, parsear_entero,
    normalizar_texto, leer_resultados, leer_alumnos, CATALOGO_PLANES,
)

ENCABEZADO_RESULTADOS = [
    "OP", "CARRERA", "MATERIA", "DOCENTE", "FECHA", "CEDULA", "NOMBRE", "ESTUDIANTE",
    "CRL", "EVAL", "PUNT", "PPRO", "RESULT.", "CRED", "OBS", "Proy", "Zona",
]
ENCABEZADO_ALUMNOS = [
    "clave", "cl_codigo", "cl_nombre", "cl_rsocial", "cl_direc", "cl_local", "cl_cpostal",
    "cl_ruc", "cl_telef", "cl_fax", "cl_email", "cl_observ", "cl_zona", "cl_obs", "cl_depto",
    "Plan", "Plan2",
]


def acta(op="ANG", plan="AP 2011", materia="PROGRAMACION 1", docente="DOCENTE, UNO",
         fecha=datetime(2014, 8, 15), cedula=41234567, nombre="PEREZ GOMEZ JUAN",
         tipo="EXA", punt=70, ppro=70, result="APR", cred="T", obs=2398, proy=None, zona=None):
    return [op, plan, materia, docente, fecha, cedula, None, nombre, None, tipo, punt, ppro,
            result, cred, obs, proy, zona]


def persona(cedula=41234567, nombre="PEREZ GOMEZ JUAN", direc="ARTIGAS 123", local="SALTO",
            telef="473 12345", fax="099 111222", email="x@y.com", zona="11AP", obs=None,
            depto="15", plan="AP", plan2=2011):
    return [cedula, f"  {cedula}", nombre, nombre, direc, local, "50000", None, telef, fax,
            email, None, zona, obs, depto, plan, plan2]


def armar_planilla(ruta, resultados, alumnos):
    wb = Workbook()
    ws = wb.active
    ws.title = "RESULTADOS"
    ws.append(ENCABEZADO_RESULTADOS)
    for fila in resultados:
        ws.append(fila)
    wa = wb.create_sheet("ALUMNOS")
    wa.append(ENCABEZADO_ALUMNOS)
    for fila in alumnos:
        wa.append(fila)
    wb.save(ruta)
    return str(ruta)


@pytest.fixture(name="planilla")
def fixture_planilla(tmp_path):
    resultados = [
        acta(),                                                       # fila 2
        acta(),                                                       # 3: duplicado exacto
        acta(materia="PROGRAMACION 1 ", result="EXO "),               # 4: distinta (espacios: se normaliza)
        acta(plan="AP2011", materia="BASES DE DATOS 1", tipo="CUR",
             punt=87, ppro=87, result="EXO", cred="T", obs=2404),     # 5: plan sin espacio
        acta(plan="OPDG", materia="PHOTOSHOP CTC", tipo="TALLER",
             cedula=41000001, nombre="NÚÑEZ PÉREZ ANA", obs="1351 B"),  # 6: alias OPDG, acta texto, no esta en ALUMNOS
        acta(fecha="31/07//2024", materia="PENSAMIENTO COMPUTACIONAL",
             tipo="CUR", punt=83, ppro=83, result="APR", cred="P"),  # 7: fecha tipeada rescatable
        acta(fecha="31/11/2024", materia="EXAMEN INTEGRADOR"),        # 8: fecha imposible
        acta(materia="INFORMATICA 1", result="#REF!"),                # 9: basura en RESULT
        acta(materia="INTRODUCCION AL COMPUTADOR", tipo="CUR",
             punt="REV", ppro=0, result="REV", obs=None),             # 10: PUNT no numerico
        ["YG"] + [None] * 16,                                         # 11: solo operador
        acta(plan="PLAN RARO", materia="ALGO"),                       # 12: plan que no esta en el catalogo
        acta(cedula="41234568", materia="TALLER DE TECNOLOGIAS", tipo="TALLER"),  # 13: cedula como texto
    ]
    alumnos = [
        persona(),                                                    # 2
        persona(direc=None, telef=None, email=None),                  # 3: repetida con menos datos
        persona(cedula=5, nombre="COBRANZA A IMPUTAR"),               # 4: cuenta contable
        persona(cedula="#REF!", nombre="ROTO"),                       # 5
        persona(cedula=30000009, nombre="LOPEZ SILVA NESTOR", obs=True, depto=True,
                plan=True, plan2="87"),                               # 6: fila desplazada
        persona(cedula=30000010, nombre="SIN ACTAS JUAN", email="099 862756",
                plan="otro plan", plan2=2001, zona="1"),              # 7: email que es telefono
    ]
    return armar_planilla(tmp_path / "escolaridades.xlsx", resultados, alumnos)


# ── Normalizacion ────────────────────────────────────────────────────────────

class TestNormalizacion:
    def test_codigo_plan(self):
        assert codigo_plan("AP2011") == "AP 2011"
        assert codigo_plan(" ap 2011 ") == "AP 2011"
        assert codigo_plan("AP2020") == "AP 2020"
        assert codigo_plan("OPDG") == "OP DG"
        assert codigo_plan("TA 2016 N") == "TA 2016 N"
        assert codigo_plan("At.C") == "AT.C"
        assert codigo_plan("DG ") == "DG"
        assert codigo_plan(None) is None

    def test_todos_los_codigos_de_la_planilla_estan_en_el_catalogo(self):
        """Los 42 que aparecen en RESULTADOS, ya normalizados."""
        vistos = [
            "TA 2001", "AP 2011", "AP 2007", "SJ 2001", "AP 2022", "TA 2016", "AP2020", "TSI 2011",
            "TA 2016 N", "OP DG", "OPC", "TGDE 2022", "TEI", "OP CONT", "OPDG", "EXCEL AV", "TA 2011",
            "AP 2004", "TEI 2006", "AC 2016", "MANT PC", "TA 2001 N", "OPG", "D WEB", "MANT", "LOG 2016",
            "AP2011", "EX+PBI", "XP", "DG ", "LS+GNS", "AP2007", "RRHH", "ST IT", "MD", "CYN", "CDAv",
            "AC", "At.C", "TA 2017", "LS", "IA 2024",
        ]
        faltan = {codigo_plan(v) for v in vistos} - set(CATALOGO_PLANES)
        assert not faltan

    def test_parsear_fecha(self):
        assert parsear_fecha(datetime(2014, 8, 15)) == (date(2014, 8, 15), None)
        assert parsear_fecha("31/07//2024") == (date(2024, 7, 31), None)
        assert parsear_fecha("14/032025") == (date(2025, 3, 14), None)
        assert parsear_fecha("12/12//2025") == (date(2025, 12, 12), None)
        assert parsear_fecha("31/11/2024") == (None, "31/11/2024")
        assert parsear_fecha(None) == (None, None)

    def test_parsear_cedula(self):
        assert parsear_cedula(41234567) == "41234567"
        assert parsear_cedula("  41234567") == "41234567"
        assert parsear_cedula(41234567.0) == "41234567"
        assert parsear_cedula(5) is None
        assert parsear_cedula("#REF!") is None
        assert parsear_cedula(211536100010) is None   # RUC

    def test_parsear_entero(self):
        assert parsear_entero(87) == 87
        assert parsear_entero(87.0) == 87
        assert parsear_entero("REV") is None
        assert parsear_entero(True) is None

    def test_normalizar_texto(self):
        assert normalizar_texto("  Programacion   1 ") == "PROGRAMACION 1"
        assert normalizar_texto("PESTAÐA") == "PESTAÑA"
        assert normalizar_texto(True) is None
        assert normalizar_texto("x" * 200, 120) == "X" * 120


# ── Lectura ──────────────────────────────────────────────────────────────────

class TestLeerResultados:
    def test_lo_que_se_carga_y_lo_que_se_descarta(self, planilla):
        from openpyxl import load_workbook
        ws = load_workbook(planilla, read_only=True, data_only=True)["RESULTADOS"]
        filas, stats = leer_resultados(ws)

        assert stats["duplicadas"] == 1
        assert stats["vacias"] == 1
        assert stats["fecha_ilegible"] == 1
        assert stats["cargadas"] == 10
        assert [f["fila_origen"] for f in filas] == [2, 4, 5, 6, 7, 8, 9, 10, 12, 13]

    def test_normalizaciones_fila_por_fila(self, planilla):
        from openpyxl import load_workbook
        ws = load_workbook(planilla, read_only=True, data_only=True)["RESULTADOS"]
        filas = {f["fila_origen"]: f for f in leer_resultados(ws)[0]}

        assert filas[4]["materia"] == "PROGRAMACION 1" and filas[4]["resultado"] == "EXO"
        assert filas[5]["plan"] == "AP 2011"
        assert filas[6]["plan"] == "OP DG" and filas[6]["acta"] == "1351 B"
        assert filas[7]["fecha"] == date(2024, 7, 31) and filas[7]["fecha_texto"] is None
        assert filas[8]["fecha"] is None and filas[8]["fecha_texto"] == "31/11/2024"
        assert filas[9]["resultado"] is None
        assert filas[10]["puntaje"] is None and filas[10]["puntaje_promedio"] == 0 and filas[10]["acta"] is None
        assert filas[13]["cedula"] == "41234568"
        assert filas[2]["acta"] == "2398" and filas[2]["operador"] == "ANG"

    def test_encabezado_distinto_frena(self, tmp_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "RESULTADOS"
        ws.append(["OP", "CARRERA", "ASIGNATURA"])
        wb.create_sheet("ALUMNOS").append(ENCABEZADO_ALUMNOS)
        ruta = tmp_path / "otra.xlsx"
        wb.save(ruta)
        from openpyxl import load_workbook
        with pytest.raises(ValueError, match="MATERIA"):
            leer_resultados(load_workbook(ruta, read_only=True)["RESULTADOS"])


class TestLeerAlumnos:
    def test_personas_y_descartes(self, planilla):
        from openpyxl import load_workbook
        ws = load_workbook(planilla, read_only=True, data_only=True)["ALUMNOS"]
        personas, stats = leer_alumnos(ws)

        assert stats["personas"] == 3
        assert stats["cedulas_repetidas"] == 1
        assert stats["desplazadas"] == 1
        assert stats["sin_cedula_valida"] == 2

    def test_se_queda_la_fila_con_mas_datos(self, planilla):
        from openpyxl import load_workbook
        ws = load_workbook(planilla, read_only=True, data_only=True)["ALUMNOS"]
        personas, _ = leer_alumnos(ws)
        p = personas["41234567"]
        assert p["fila_origen"] == 2 and p["direccion"] == "ARTIGAS 123"
        assert p["plan_declarado"] == "AP 2011" and p["zona"] == "11AP"
        assert p["celular"] == "099 111222" and p["email"] == "x@y.com"

    def test_fila_desplazada_solo_cedula_y_nombre(self, planilla):
        from openpyxl import load_workbook
        ws = load_workbook(planilla, read_only=True, data_only=True)["ALUMNOS"]
        p = leer_alumnos(ws)[0]["30000009"]
        assert p["nombre"] == "LOPEZ SILVA NESTOR"
        assert "direccion" not in p and "desplazada" in p["observaciones"]

    def test_otro_plan_y_zona_suelta_quedan_en_null(self, planilla):
        from openpyxl import load_workbook
        ws = load_workbook(planilla, read_only=True, data_only=True)["ALUMNOS"]
        p = leer_alumnos(ws)[0]["30000010"]
        assert p["plan_declarado"] is None and p["zona"] is None
        assert p["email"] is None   # era un telefono


# ── Carga completa ───────────────────────────────────────────────────────────

class TestImportar:
    def test_carga_todo(self, session, planilla):
        informe = importar(planilla, session)

        assert informe["carga"]["planes"] == 3          # AP 2011, OP DG, PLAN RARO
        assert informe["carga"]["planes_sin_catalogo"] == 1
        assert informe["carga"]["alumnos"] == 5         # 3 de ALUMNOS + 2 solo en actas
        assert informe["carga"]["alumnos_solo_en_actas"] == 2
        assert informe["carga"]["resultados"] == 10

        raro = session.exec(select(HistoricoPlan).where(HistoricoPlan.codigo == "PLAN RARO")).one()
        assert raro.carrera is None
        ap = session.exec(select(HistoricoPlan).where(HistoricoPlan.codigo == "AP 2011")).one()
        assert (ap.carrera, ap.creditos_requeridos) == ("Analista Programador", 15)

    def test_quien_solo_esta_en_actas_toma_el_nombre_del_acta(self, session, planilla):
        importar(planilla, session)
        ana = session.exec(select(HistoricoAlumno).where(HistoricoAlumno.cedula == "41000001")).one()
        assert ana.nombre == "NÚÑEZ PÉREZ ANA"
        assert ana.nombre_busqueda == "NUNEZ PEREZ ANA"
        assert "acta" in ana.observaciones

    def test_las_filas_apuntan_a_su_alumno_y_plan(self, session, planilla):
        importar(planilla, session)
        fila = session.exec(select(HistoricoResultado).where(HistoricoResultado.fila_origen == 6)).one()
        assert fila.alumno.cedula == "41000001" and fila.plan.codigo == "OP DG"

    def test_no_pisa_sin_reemplazar(self, session, planilla):
        importar(planilla, session)
        with pytest.raises(ValueError, match="reemplazar"):
            importar(planilla, session)
        assert session.exec(select(func.count(HistoricoResultado.id))).one() == 10

    def test_reemplazar_deja_una_sola_copia(self, session, planilla):
        importar(planilla, session)
        importar(planilla, session, reemplazar=True)
        assert session.exec(select(func.count(HistoricoResultado.id))).one() == 10
        assert session.exec(select(func.count(HistoricoAlumno.id))).one() == 5

    def test_dry_run_no_escribe(self, session, planilla):
        informe = importar(planilla, session, dry_run=True)
        assert informe["carga"]["resultados"] == 10
        assert session.exec(select(func.count(HistoricoResultado.id))).one() == 0

    def test_planilla_sin_las_hojas(self, session, tmp_path):
        wb = Workbook()
        wb.active.title = "OTRA"
        ruta = tmp_path / "x.xlsx"
        wb.save(ruta)
        with pytest.raises(ValueError, match="RESULTADOS"):
            importar(str(ruta), session)
